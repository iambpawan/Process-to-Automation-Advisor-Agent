"""Design-document assembly. All numbers are rendered here, never by the LLM."""

from __future__ import annotations

import io
import re

from docx import Document

from advisor.agents.specialists import Assessment

PLACEHOLDERS = (
    "PROCESS_TABLE", "FIT_SCORE", "AI_DECISION", "DELIVERY", "EXISTING_AUTOMATION", "COMPLIANCE", "COST_TABLE",
)
MAX_FLOW_NAMES = 10


def _money(v: float, cur: str) -> str:
    return f"{v:,.0f} {cur}"


def _table(headers: list[str], rows: list[list[str]]) -> str:
    lines = ["| " + " | ".join(headers) + " |", "|" + "---|" * len(headers)]
    lines += ["| " + " | ".join(str(c).replace("|", "/") for c in r) + " |" for r in rows]
    return "\n".join(lines)


def _yn(v: bool | None) -> str:
    return "—" if v is None else ("Yes" if v else "No")


def process_table(a: Assessment) -> str:
    p = a.profile
    head = (
        f"**Volume:** {p.transactions_per_month:,.0f} transactions/month · "
        f"**Handling time:** {p.total_minutes():.0f} min/transaction · "
        f"**Manual effort:** {a.costs.current_hours_monthly:,.0f} h/month "
        f"({_money(a.costs.current_monthly_cost, a.currency)}/month)\n\n"
    )
    rows = [
        [s.name, f"{s.minutes_per_run or 0:g}", s.system or "—", (s.input_type or "—").replace("_", " "),
         _yn(s.requires_judgment), f"{s.exception_rate_pct or 0:g}%"]
        for s in p.steps
    ]
    return head + _table(["Step", "Min", "System", "Input", "Judgment", "Exceptions"], rows)


def fit_score(a: Assessment) -> str:
    f = a.fit
    rows = [[c.name, f"{c.points:g} / {c.max_points}", c.note] for c in f.components]
    return f"**Automation fit score: {f.total} / 100 — {f.band}**\n\n" + _table(["Factor", "Points", "Basis"], rows)


def ai_decision(a: Assessment) -> str:
    d, c = a.ai_decision, a.costs
    out = [f"**Recommendation: {d.approach}** — {d.reason}"]
    if d.alternative:
        out.append(f"\n*Alternative:* {d.alternative}")
    if d.ai_steps:
        out.append("\n**AI signals found**\n")
        out.append(_table(["Step", "Why AI is needed"], [[s.step, ", ".join(s.reasons)] for s in d.ai_steps]))
        out.append(
            f"\n**AI cost check:** {c.ai_cost_per_transaction:,.4f} {c.currency} per transaction × "
            f"{a.profile.transactions_per_month:,.0f}/month = {_money(c.ai_monthly_cost, c.currency)}/month, "
            f"against {_money(c.ai_steps_labor_saved_monthly, c.currency)}/month of labour saved on those steps."
        )
        out.append("\n**Benefits of AI here**\n" + "\n".join(f"- {b}" for b in d.benefits))
        out.append("\n**Limitations**\n" + "\n".join(f"- {x}" for x in d.limitations))
    return "\n".join(out)


def delivery_section(a: Assessment) -> str:
    rows = [[p.step, p.tool, p.human_reason or "—"] for p in a.step_plans]
    e, m = a.effort, a.maintenance
    return "\n".join(
        [
            _table(["Step", "Recommended tool", "Human in the loop"], rows),
            "\n**Platform components:** " + "; ".join(a.platforms),
            f"\n**Effort:** {e.size} — about {e.days:g} build days (range {e.low_days:g}–{e.high_days:g}).",
            f"\n**Maintenance risk:** {m.risk} — " + "; ".join(m.drivers),
        ]
    )


def compliance_section(a: Assessment) -> str:
    c = a.compliance
    lines = [
        f"**Sensitive data:** {', '.join(c.categories) or 'None declared'}",
        f"\n**Can data be sent to the LLM?** {'Yes, with conditions' if c.can_send_to_llm and c.sensitive else 'Yes' if c.can_send_to_llm else 'No'} — {c.llm_condition}",
    ]
    if c.regulations:
        lines.append("\n**Regulations to check**\n" + "\n".join(f"- {r}" for r in c.regulations))
    if c.required_controls:
        lines.append("\n**Required controls**\n" + "\n".join(f"- {r}" for r in c.required_controls))
    return "\n".join(lines)


def cost_table(a: Assessment) -> str:
    c, cur = a.costs, a.currency

    def payback(s):
        return "never" if s.payback_months is None else f"{s.payback_months:.1f} months"

    def roi(s):
        return "—" if s.roi_pct is None else f"{s.roi_pct:,.0f}%"

    rows = [
        ["Steps automated", ", ".join(c.rules_only.automated_steps) or "—", ", ".join(c.with_ai.automated_steps)],
        ["Labour saved / month", _money(c.rules_only.labor_saved_monthly, cur), _money(c.with_ai.labor_saved_monthly, cur)],
        ["Power Automate licences / month", _money(c.rules_only.pa_license_monthly, cur), _money(c.with_ai.pa_license_monthly, cur)],
        ["Unattended bots", str(c.rules_only.bots), str(c.with_ai.bots)],
        ["AI usage / month", _money(c.rules_only.ai_monthly, cur), _money(c.with_ai.ai_monthly, cur)],
        ["Support & maintenance / month", _money(c.rules_only.support_monthly, cur), _money(c.with_ai.support_monthly, cur)],
        ["Net saving / month", _money(c.rules_only.net_monthly, cur), _money(c.with_ai.net_monthly, cur)],
        ["Build cost (one-time)", _money(c.rules_only.build_cost, cur), _money(c.with_ai.build_cost, cur)],
        ["Payback", payback(c.rules_only), payback(c.with_ai)],
        [f"{c.horizon_years}-year net value", _money(c.rules_only.horizon_value, cur), _money(c.with_ai.horizon_value, cur)],
        [f"{c.horizon_years}-year ROI", roi(c.rules_only), roi(c.with_ai)],
    ]
    chosen = "Rules + AI" if a.ai_decision.ai_needed else "Rules-only"
    return (
        f"**Current manual cost:** {_money(c.current_monthly_cost, cur)}/month. **Recommended scenario:** {chosen}.\n\n"
        + _table(["", "Rules-only", "Rules + AI"], rows)
        + "\n\n*Assumptions:* " + "; ".join(c.assumptions)
    )


def existing_automation(a: Assessment) -> str:
    """Existing flows that may already cover steps, and licence cost avoided by reusing spare licences."""
    out = []
    if a.flows:
        names = ", ".join(f.name for f in a.flows[:MAX_FLOW_NAMES]) + ("…" if len(a.flows) > MAX_FLOW_NAMES else "")
        out.append(f"**Existing flows reviewed:** {len(a.flows)} ({names})")
        if a.overlaps:
            rows = [[o.step, o.flow, f"{o.score:.0%}", o.reason] for o in a.overlaps]
            out.append("\n" + _table(["Process step", "Existing flow", "Match", "Why"], rows))
            out.append("\n*Matches are word-based suggestions: check each flow before reusing or extending it.*")
        else:
            out.append("\nNo existing flow appears to cover a process step.")
    else:
        out.append("No existing flows provided.")
    inv, c = a.inventory, a.costs
    if inv is None:
        out.append("\nNo licence inventory provided, so the full licence cost is counted.")
    else:
        out.append(
            f"\n**Spare licences in tenant:** {inv.premium_available} Premium, {inv.process_available} Process, "
            f"{inv.ai_builder_credits_available:,} AI Builder credits/month"
        )
        rows = [
            [s.name, _money(s.license_offset_monthly, a.currency), "; ".join(s.offset_notes) or "none apply"]
            for s in (c.rules_only, c.with_ai)
        ]
        out.append("\n" + _table(["Scenario", "Licence cost avoided / month", "Reused"], rows))
        out.append("\n*Assumes these spare licences are not already earmarked for other projects.*")
    return "\n".join(out)


RENDERERS = {
    "PROCESS_TABLE": process_table,
    "FIT_SCORE": fit_score,
    "AI_DECISION": ai_decision,
    "DELIVERY": delivery_section,
    "EXISTING_AUTOMATION": existing_automation,
    "COMPLIANCE": compliance_section,
    "COST_TABLE": cost_table,
}


def fill(narrative: str, a: Assessment) -> str:
    """Replace {{PLACEHOLDER}} tokens with Python-rendered sections; append any the writer left out."""
    doc = narrative
    missing = []
    for key, render in RENDERERS.items():
        # tolerate "{{ cost_table }}" and other case/space variants a model may write
        token = re.compile(r"\{\{\s*" + key + r"\s*\}\}", re.I)
        if token.search(doc):
            rendered = render(a)
            doc = token.sub(lambda _m: rendered, doc)
        else:
            missing.append(render(a))
    doc = re.sub(r"\{\{\s*[A-Za-z_]+\s*\}\}", "", doc)  # drop unknown leftover tokens
    if missing:
        doc += "\n\n## Appendix: Analysis details\n\n" + "\n\n".join(missing)
    return doc


_TABLE_SEP = re.compile(r"^\|(\s*:?-+:?\s*\|)+\s*$")
# Characters XML 1.0 (and so Word) cannot store; pasted PDF/Word text often carries them.
_XML_ILLEGAL = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f￾￿]")


def strip_control_chars(text: str) -> str:
    """Remove characters that are illegal in XML 1.0."""
    return _XML_ILLEGAL.sub("", text)


def to_docx(markdown: str) -> bytes:
    """Minimal Markdown -> DOCX: headings, bullets, tables, paragraphs (bold markers stripped)."""
    doc = Document()
    lines = markdown.splitlines()
    i = 0

    def clean(text: str) -> str:
        return strip_control_chars(text).replace("**", "").replace("*", "").strip()

    while i < len(lines):
        line = lines[i].rstrip()
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if not _TABLE_SEP.match(lines[i].strip()):
                    rows.append([clean(c) for c in lines[i].strip().strip("|").split("|")])
                i += 1
            if not rows:
                continue  # only separator rows: nothing to draw
            table = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
            table.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, cell in enumerate(row):
                    table.cell(r, c).text = cell
            continue
        heading = re.match(r"^(#{1,4})\s+(.*)", line)
        if heading:
            doc.add_heading(clean(heading.group(2)), level=len(heading.group(1)) - 1 or 0)
        elif re.match(r"^\s*[-*]\s+", line):
            doc.add_paragraph(clean(re.sub(r"^\s*[-*]\s+", "", line)), style="List Bullet")
        elif line.strip():
            doc.add_paragraph(clean(line))
        i += 1
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
